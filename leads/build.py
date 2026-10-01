"""Merge verification results into the original workbook.

Writes: Master_Leads_-_Best_600_verified.xlsx (same sheet/columns, values updated,
plus 3 audit columns at the end).  Changed cells are highlighted.
"""
import json, glob, os, re
import openpyxl
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'Master_Leads_-_Best_600_original.xlsx')
OUT = os.path.join(HERE, 'Master_Leads_-_Best_600_verified.xlsx')

# ---- gather results from main file + agent slice files
res = {}
for p in [os.path.join(HERE, 'results.json')] + sorted(glob.glob(os.path.join(HERE, 'results_[A-H].json'))):
    if os.path.exists(p):
        for k, v in json.load(open(p)).items():
            res.setdefault(k, {}).update(v)

CONFIRMED_RE = re.compile(r'(confirmed|found on their own site|found via search|found in director|found via director|'
                          r'email corrected|email fixed|email found|branch email|brand site|own site|directory listing|'
                          r'directory email|customer-care email|general email|group email|hmc customer)', re.I)
UNVERIFIED_RE = re.compile(r'(unverified|kept|obfuscated|pending)', re.I)
UNSEARCHED_RE = re.compile(r'(NOT SEARCHED|no live search|budget exhausted|budget was exhausted|search budget|needs re-run)', re.I)

def classify(r, old_email):
    email = (r.get('email') or '').strip()
    st = r.get('status', '') or ''
    if not email:
        return 'No email found'
    if email.lower() == old_email.lower():
        if CONFIRMED_RE.search(st) and not UNVERIFIED_RE.search(st):
            return 'Confirmed'
        if UNVERIFIED_RE.search(st):
            return 'Kept, unverified'
        return 'Confirmed' if CONFIRMED_RE.search(st) else 'Kept, unverified'
    # email changed
    if UNVERIFIED_RE.search(st) and not CONFIRMED_RE.search(st):
        return 'Corrected, unverified'
    return 'Corrected'

wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[0]
hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}
col = lambda name: hdr[name]

# audit columns
new_cols = ['Email status', 'Previous email', 'Previous website']
for name in new_cols:
    if name not in hdr:
        ws.cell(row=1, column=ws.max_column + 1, value=name)
        hdr[name] = ws.max_column
# copy header style from column A
h0 = ws.cell(row=1, column=1)
for name in new_cols:
    c = ws.cell(row=1, column=hdr[name])
    c.font = copy(h0.font); c.fill = copy(h0.fill); c.alignment = copy(h0.alignment); c.border = copy(h0.border)

FILL_EMAIL_CHANGED = PatternFill('solid', fgColor='FFF2CC')   # light yellow
FILL_WEB_CHANGED = PatternFill('solid', fgColor='FCE4D6')     # light orange
FILL_NONE = PatternFill('solid', fgColor='F8CBAD')            # light red
FILL_UNVERIFIED = PatternFill('solid', fgColor='EDEDED')      # grey

stats = {'rows': 0, 'processed': 0, 'Confirmed': 0, 'Corrected': 0, 'Corrected, unverified': 0,
         'Kept, unverified': 0, 'No email found': 0, 'website_changed': 0, 'not_processed': 0}

for row in range(2, ws.max_row + 1):
    n = ws.cell(row=row, column=col('#')).value
    if n is None:
        continue
    stats['rows'] += 1
    r = res.get(str(int(n)))
    old_email = (ws.cell(row=row, column=col('Email')).value or '').strip()
    old_web = (ws.cell(row=row, column=col('Website')).value or '').strip()
    if not r or r.get('status') == 'pending followup':
        stats['not_processed'] += 1
        ws.cell(row=row, column=col('Email status')).value = 'Not yet searched (session search cap)'
        ws.cell(row=row, column=col('Check')).value = 'NOT SEARCHED: session web-search limit reached before this row - needs re-run'
        continue
    stats['processed'] += 1
    email = (r.get('email') or '').strip()
    status_code = classify(r, old_email)
    if UNSEARCHED_RE.search(r.get('status') or ''):
        status_code = 'Not yet searched (session search cap)'
    stats[status_code] = stats.get(status_code, 0) + 1

    # Email
    c = ws.cell(row=row, column=col('Email'))
    if email.lower() != old_email.lower():
        c.value = email if email else None
        c.fill = FILL_EMAIL_CHANGED if email else FILL_NONE
        ws.cell(row=row, column=col('Previous email')).value = old_email or None
    elif status_code == 'Kept, unverified':
        c.fill = FILL_UNVERIFIED
    # Email (confirmed)
    ws.cell(row=row, column=col('Email (confirmed)')).value = email if status_code in ('Confirmed', 'Corrected') else None
    # All emails found
    ws.cell(row=row, column=col('All emails found')).value = (r.get('all') or email or None)
    # Where it came from
    ws.cell(row=row, column=col('Where it came from')).value = (r.get('source') or None)
    # Check
    ws.cell(row=row, column=col('Check')).value = (r.get('status') or None)
    # Website
    if 'website' in r and r['website'] is not None and r['website'].strip() != old_web:
        wc = ws.cell(row=row, column=col('Website'))
        wc.value = r['website'].strip() or None
        wc.fill = FILL_WEB_CHANGED
        ws.cell(row=row, column=col('Previous website')).value = old_web or None
        stats['website_changed'] += 1
    # status code
    ws.cell(row=row, column=col('Email status')).value = status_code

# widths for new columns
for name, w in (('Email status', 22), ('Previous email', 34), ('Previous website', 44), ('Check', 70)):
    ws.column_dimensions[openpyxl.utils.get_column_letter(hdr[name])].width = w

# legend sheet
lg = wb.create_sheet('Legend')
rows = [
    ('Colour / value', 'Meaning'),
    ('Yellow fill on Email', 'Email was changed - the previous value is in "Previous email"'),
    ('Red fill on Email', 'No valid email could be found - previous (junk) value is in "Previous email"'),
    ('Grey fill on Email', 'Held email kept but could not be confirmed by search (domain matches their site)'),
    ('Orange fill on Website', 'Website was corrected - the previous value is in "Previous website"'),
    ('Email status = Confirmed', 'Email seen on the company\'s own site or a reputable directory'),
    ('Email status = Corrected', 'Previous email was wrong/junk; a confirmed replacement was found'),
    ('Email status = Corrected, unverified', 'Previous email was wrong; replacement found but not fully confirmed'),
    ('Email status = Kept, unverified', 'Could not confirm; kept the existing address'),
    ('Email status = No email found', 'No usable email found by search - phone/social noted in "Check"'),
    ('Email status = Not yet searched (session search cap)', 'Row was NOT verified online: the session hit its web-search limit. Only rule-based cleanup was applied (obvious placeholder/junk emails removed). Needs a re-run.'),
    ('Check column', 'What was found and where; includes phone numbers where seen'),
    ('Method', 'Verified by web search (company sites could not be fetched directly from this environment).'),
]
for i, (a, b) in enumerate(rows, 1):
    lg.cell(row=i, column=1, value=a); lg.cell(row=i, column=2, value=b)
lg['A1'].font = Font(bold=True); lg['B1'].font = Font(bold=True)
lg.column_dimensions['A'].width = 38; lg.column_dimensions['B'].width = 95

wb.save(OUT)
print(json.dumps(stats, indent=1))
print('saved', OUT)
